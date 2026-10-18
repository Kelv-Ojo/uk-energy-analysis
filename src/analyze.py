"""Rebuild a descriptive UK energy case study from the original DESNZ workbook."""
from pathlib import Path
import re
import json
import sqlite3
import openpyxl
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]

def main():
    out = ROOT/'reports'
    out.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(ROOT/'data/ECUK_2025_Consumption_tables.xlsx', read_only=True, data_only=True)
    sheet = wb['Table C1']
    assert 'ktoe' in sheet['A1'].value
    rows = list(sheet.iter_rows(min_row=5, max_row=61, max_col=57, values_only=True))
    sectors = {0: ('Industry', 12), 12: ('Transport', 26), 26: ('Domestic', 36),
               36: ('Services', 45), 45: ('All final users', 57)}
    records = []
    for start, (sector, end) in sectors.items():
        for row in rows[2:]:
            year = row[start]
            if not isinstance(year, (int, float)):
                continue
            for col in range(start+1, end):
                label = re.sub(r'\s*\[Note \d+\]', '', rows[1][col]).replace('\n',' ').strip()
                raw = row[col]
                value = float(raw) if isinstance(raw, (int, float)) else None
                records.append(dict(year=int(year), sector=sector, fuel=label,
                                    consumption_ktoe=value, consumption_twh=None if value is None else value*0.01163,
                                    status='reported' if value is not None else 'unavailable',
                                    source_cell=f'{openpyxl.utils.get_column_letter(col+1)}{int(year)-1970+7}'))
    data = pd.DataFrame(records)
    assert not data.duplicated(['year','sector','fuel']).any()
    assert data.year.min()==1970 and data.year.max()==2024
    assert data.consumption_ktoe.dropna().ge(0).all()
    totals = data[data.fuel.eq('Total')].pivot(index='year', columns='sector', values='consumption_ktoe')
    parts = ['Industry','Transport','Domestic','Services']
    assert totals.shape==(55,5) and totals.notna().all().all()
    residual = totals[parts].sum(axis=1)-totals['All final users']
    # Independently published totals are rounded; retain and report the discrepancy.
    assert residual.abs().max()<10
    data.to_csv(out/'energy_tidy.csv',index=False)
    totals.to_csv(out/'sector_totals_ktoe.csv')
    with sqlite3.connect(':memory:') as db:
        data.to_sql('energy',db,index=False)
        query=(ROOT/'sql/sector_changes.sql').read_text()
        summary=pd.read_sql_query(query,db)
        summary.to_csv(out/'sector_changes.csv',index=False)
    plt.rcParams.update({'font.family':'DejaVu Sans','axes.spines.top':False,'axes.spines.right':False})
    fig,ax=plt.subplots(figsize=(10,5.5))
    (totals[parts]*0.01163).plot(ax=ax,linewidth=2)
    ax.set(title='UK final energy consumption by sector, 1970–2024',xlabel='Year',ylabel='TWh')
    ax.legend(frameon=False,ncol=2)
    fig.tight_layout();fig.savefig(out/'sector_trends.png',dpi=160);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,5))
    changes=((totals[parts].iloc[-1]/totals[parts].iloc[0]-1)*100).sort_values()
    changes.plot.barh(ax=ax,color=['#176B87' if x<0 else '#D87A36' for x in changes])
    ax.axvline(0,color='#333',linewidth=.8)
    ax.set(title='Change in final energy consumption, 1970–2024',xlabel='Change (%)',ylabel='')
    fig.tight_layout();fig.savefig(out/'sector_changes.png',dpi=160);plt.close(fig)
    total_change=(totals['All final users'].iloc[-1]/totals['All final users'].iloc[0]-1)*100
    yoy=(totals['All final users'].iloc[-1]/totals['All final users'].iloc[-2]-1)*100
    transport_share=totals.Transport.iloc[-1]/totals['All final users'].iloc[-1]*100
    metrics={'years':55,'first_year':1970,'last_year':2024,'total_change_pct':total_change,
             'latest_yoy_pct':yoy,'transport_share_pct':transport_share,
             'missing_component_values':int(data.consumption_ktoe.isna().sum()),
             'max_sector_reconciliation_difference_ktoe':float(residual.abs().max())}
    (out/'metrics.json').write_text(json.dumps(metrics,indent=2)+'\n')
    lines=['# UK energy: findings and decision context','',
           f'- Final energy consumption changed by **{total_change:.1f}%** between 1970 and 2024.',
           f'- Consumption rose **{yoy:.1f}%** between 2023 and 2024.',
           f'- Transport accounts for **{transport_share:.1f}%** of reported final consumption in 2024.','',
           '## Suggested next steps','',
           'Prioritize transport for deeper demand analysis because it represents a large share of final consumption. '+
           'Investigate activity, weather and efficiency data before explaining year-to-year changes or setting targets.','',
           '## Method and limits','',
           'Table C1 is in ktoe. Conversion: 1 ktoe = 0.01163 TWh. Use source totals; never add an aggregate sector to its components. '+
           'Unavailable fuel cells remain missing rather than becoming zero. Differences in published sector sums are retained as rounding diagnostics. '+
           'This is a descriptive historical analysis of the local 2025 workbook snapshot, not a current release, causal study, emissions inventory or forecast. '+
           'Electricity is a delivered fuel and cannot be labelled renewable from this table.','',
           f'Validation: 55 consecutive years; unique year/sector/fuel keys; nonnegative reported values; complete sector totals. '+
           f'Maximum reconciliation difference: {residual.abs().max():.2f} ktoe.']
    (out/'findings.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(metrics,indent=2))

if __name__=='__main__':
    main()
